"""Leitura de preço: insumo (ISD/ICD/ISE) e composição (CSD/CCD/CSE).

Porquê as duas abas num arquivo só: são o mesmo tipo de leitura —
preço por UF, num regime — com o único layout que muda de verdade
entre elas (§2.1). Separar em dois arquivos duplicaria o cabeçalho
sem separar responsabilidade nenhuma.
"""

from __future__ import annotations

from collections.abc import Iterator
from typing import NamedTuple, cast

import openpyxl

from core.reference_base.adapters.sinapi.spreadsheet import (
    localizar_linha_cabecalho,
    mapa_ufs_por_coluna_unica,
    mapa_ufs_por_par_de_colunas,
    normalizar,
    valor_para_centavos,
    valor_para_texto_decimal,
)
from core.reference_base.constants import (
    REGIME_ONERADO,
    ROTULO_CABECALHO_COMPOSICAO,
    ROTULO_CABECALHO_INSUMO,
    TIPO_COMPOSICAO,
    TIPO_INSUMO,
    UFS_VALIDAS,
)
from core.reference_base.literals import Regime, TipoItem
from core.reference_base.types import RegistroItem, RegistroPreco

NUMERO_COLUNAS_INSUMO: int = 5  # Classificação, Código, Descrição, Unidade, Origem


class ChaveDeRecuperacaoAmbigua(ValueError):
    """V3 falhou e a chave (descrição, unidade) não resolve um código.

    Porquê: S03 dos testes de aceite exige abortar, não adivinhar.
    """


class _LinhaInsumo(NamedTuple):
    classification: object
    codigo: object
    descricao: object
    unidade: object
    origem_preco: object


class _LinhaComposicao(NamedTuple):
    """As quatro colunas fixas — o resto da linha varia com o número de UFs."""

    grupo: object
    codigo: object
    descricao: object
    unidade: object


def ler_insumos_onerados(
    livro: openpyxl.Workbook, aba: str, situacao_por_codigo: dict[str, str]
) -> tuple[list[RegistroItem], list[RegistroPreco]]:
    """Lê ISD uma vez só: monta o catálogo E o preço onerado juntos.

    Porquê fundir os dois: catálogo (código, descrição, unidade) e
    preço onerado vêm exatamente da mesma aba. Lê-la duas vezes — uma
    para cada — mediu ~6 s à toa no pacote real; ~1,25 milhão de
    linhas de preço no total não perdoa leitura repetida (§12).
    """
    planilha = livro[aba]
    linhas = list(planilha.iter_rows(values_only=True))
    indice_cabecalho = localizar_linha_cabecalho(linhas, ROTULO_CABECALHO_INSUMO)
    mapa_ufs = mapa_ufs_por_coluna_unica(linhas[indice_cabecalho], UFS_VALIDAS)
    catalogo: list[RegistroItem] = []
    precos: list[RegistroPreco] = []
    for bruta in linhas[indice_cabecalho + 1 :]:
        if bruta[0] is None:
            continue
        item, precos_do_item = _insumo_onerado_da_linha(
            bruta, mapa_ufs, situacao_por_codigo
        )
        catalogo.append(item)
        precos.extend(precos_do_item)
    return catalogo, precos


def _insumo_onerado_da_linha(
    bruta: tuple[object, ...],
    mapa_ufs: dict[int, str],
    situacao_por_codigo: dict[str, str],
) -> tuple[RegistroItem, list[RegistroPreco]]:
    """Um insumo e seus preços onerados — a linha inteira, num só lugar."""
    campos = _LinhaInsumo(*bruta[:NUMERO_COLUNAS_INSUMO])
    codigo = str(campos.codigo).strip()
    item = RegistroItem(
        codigo=codigo,
        tipo=cast(TipoItem, TIPO_INSUMO),
        grupo=str(campos.classification).strip(),
        descricao=str(campos.descricao).strip(),
        unidade=str(campos.unidade).strip(),
        situacao=situacao_por_codigo.get(codigo),
    )
    precos = [
        RegistroPreco(
            codigo=codigo,
            tipo=cast(TipoItem, TIPO_INSUMO),
            uf=uf,
            regime=cast(Regime, REGIME_ONERADO),
            valor_centavos=valor_para_centavos(bruta[coluna]),
            percentual_as=None,
        )
        for coluna, uf in mapa_ufs.items()
    ]
    return item, precos


def ler_precos_insumo(
    livro: openpyxl.Workbook, aba: str, regime: Regime
) -> Iterator[RegistroPreco]:
    """Lê o preço mediano de cada insumo, UF por UF, para um regime.

    Porquê ainda existe, com ISD já coberto por `ler_insumos_onerados`:
    ICD e ISE são abas genuinamente diferentes — desonerado e sem
    encargos não têm como vir de uma leitura já feita.
    """
    planilha = livro[aba]
    linhas = list(planilha.iter_rows(values_only=True))
    indice_cabecalho = localizar_linha_cabecalho(linhas, ROTULO_CABECALHO_INSUMO)
    mapa_ufs = mapa_ufs_por_coluna_unica(linhas[indice_cabecalho], UFS_VALIDAS)
    for bruta in linhas[indice_cabecalho + 1 :]:
        if bruta[0] is None:
            continue
        codigo = str(bruta[1]).strip()
        for coluna, uf in mapa_ufs.items():
            yield RegistroPreco(
                codigo=codigo,
                tipo=cast(TipoItem, TIPO_INSUMO),
                uf=uf,
                regime=regime,
                valor_centavos=valor_para_centavos(bruta[coluna]),
                percentual_as=None,
            )


_COLUNA_GRUPO: int = 1
_COLUNA_CODIGO: int = 2
# +1: índice de lista (0-based) -> linha do Excel (1-based); +1: pular o
# próprio cabeçalho e começar no primeiro dado.
_DESLOCAMENTO_LINHA_DE_DADO: int = 2


def codigos_brutos_da_aba(livro: openpyxl.Workbook, aba: str) -> list[object]:
    """Coluna de código crua de uma aba de composição, para V3.

    Porquê restringir a duas colunas: V3 só precisa saber se o código
    está corrompido, não o preço. Ler as 58 colunas da aba de novo só
    para isto foi ~20 s medidos à toa no pacote real — restringir a
    coluna evita reler o que `ler_precos_composicao` já vai ler.
    """
    planilha = livro[aba]
    coluna_grupo = list(
        planilha.iter_rows(
            min_col=_COLUNA_GRUPO, max_col=_COLUNA_GRUPO, values_only=True
        )
    )
    indice_cabecalho = localizar_linha_cabecalho(
        coluna_grupo, ROTULO_CABECALHO_COMPOSICAO
    )
    linhas = planilha.iter_rows(
        min_row=indice_cabecalho + _DESLOCAMENTO_LINHA_DE_DADO,
        min_col=_COLUNA_GRUPO,
        max_col=_COLUNA_CODIGO,
        values_only=True,
    )
    return [codigo for grupo, codigo in linhas if grupo is not None]


def ler_precos_composicao(
    livro: openpyxl.Workbook,
    aba: str,
    regime: Regime,
    recuperacao_por_chave: dict[tuple[str, str], str],
) -> Iterator[RegistroPreco]:
    """Lê o custo de cada composição, UF por UF, com `%AS`, para um regime."""
    planilha = livro[aba]
    linhas = list(planilha.iter_rows(values_only=True))
    indice_cabecalho = localizar_linha_cabecalho(linhas, ROTULO_CABECALHO_COMPOSICAO)
    linha_siglas = linhas[indice_cabecalho - 1]
    mapa_ufs = mapa_ufs_por_par_de_colunas(linha_siglas, UFS_VALIDAS)
    for linha in linhas[indice_cabecalho + 1 :]:
        if linha[0] is None:
            continue
        codigo = _resolver_codigo(linha, recuperacao_por_chave)
        yield from _precos_da_linha(linha, codigo, regime, mapa_ufs)


def _resolver_codigo(
    linha: tuple[object, ...], recuperacao_por_chave: dict[tuple[str, str], str]
) -> str:
    """Usa o código bruto quando confiável; senão, recupera pela chave.

    Porquê o `float` aqui é legítimo (isenção de VD-01 em
    `adaptadores/`, VERTICE-LINT-SUITE.md §3.1): o `.xlsx` guarda todo
    número como IEEE754 — reconhecer isso não é computar com float,
    é a única forma de saber que a célula não é uma fórmula de texto.
    """
    campos = _LinhaComposicao(*linha[: len(_LinhaComposicao._fields)])
    bruto = campos.codigo
    if isinstance(bruto, int | float) and bruto not in (0, 0.0):
        return str(int(bruto))
    chave = (
        normalizar(str(campos.descricao or "")),
        normalizar(str(campos.unidade or "")),
    )
    codigo = recuperacao_por_chave.get(chave)
    if codigo is None:
        raise ChaveDeRecuperacaoAmbigua(f"sem código para a chave {chave!r}")
    return codigo


def _precos_da_linha(
    linha: tuple[object, ...],
    codigo: str,
    regime: Regime,
    mapa_ufs: dict[str, tuple[int, int]],
) -> Iterator[RegistroPreco]:
    """Emite um `RegistroPreco` por UF presente no cabeçalho da aba."""
    for uf, (coluna_custo, coluna_as) in mapa_ufs.items():
        yield RegistroPreco(
            codigo=codigo,
            tipo=cast(TipoItem, TIPO_COMPOSICAO),
            uf=uf,
            regime=regime,
            valor_centavos=valor_para_centavos(linha[coluna_custo]),
            percentual_as=valor_para_texto_decimal(linha[coluna_as]),
        )

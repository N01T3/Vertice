"""Leitores de office sem dependência obrigatória.

Porquê: DOCX e XLSX são ZIP com XML; ler o XML puro garante
funcionar offline e sem instalar nada pesado.
"""

from __future__ import annotations

import xml.etree.ElementTree as elemento_xml
import zipfile
from pathlib import Path
from typing import Any

from core.documents.base_readers import dividir_texto, origem_base
from core.documents.errors import ErroLeituraDocumento
from core.documents.types import TrechoExtraido


def extrair_texto_docx(caminho: Path) -> list[TrechoExtraido]:
    """Lê DOCX pela biblioteca quando há, senão pelo XML puro."""
    try:
        import docx as pacote_docx
    except ImportError:
        return extrair_docx_puro(caminho)
    try:
        arquivo: Any = pacote_docx.Document(str(caminho))
        paragrafos: list[str] = [
            str(paragrafo.text) for paragrafo in arquivo.paragraphs
        ]
    except OSError as erro:
        raise ErroLeituraDocumento(f"falha ao ler {caminho.name}") from erro
    conteudo = "\n".join(paragrafos)
    return dividir_texto(conteudo, "documento", origem_base(caminho))


def extrair_texto_xlsx(caminho: Path) -> list[TrechoExtraido]:
    """Lê XLSX pela biblioteca quando há, senão pelo XML puro."""
    try:
        import openpyxl as pacote_planilha
    except ImportError:
        return extrair_xlsx_puro(caminho)
    try:
        pasta: Any = pacote_planilha.load_workbook(
            str(caminho), read_only=True, data_only=True
        )
        trechos: list[TrechoExtraido] = []
        for aba in pasta.worksheets:
            trechos.extend(aba_para_trechos(aba, origem_base(caminho)))
        pasta.close()
    except OSError as erro:
        raise ErroLeituraDocumento(f"falha ao ler {caminho.name}") from erro
    return trechos


def aba_para_trechos(aba: Any, origem: str) -> list[TrechoExtraido]:
    """Converte uma aba openpyxl em trechos por linha preenchida."""
    nome_aba = str(getattr(aba, "title", "aba"))
    obtido = aba.iter_rows(values_only=True)
    saida: list[TrechoExtraido] = []
    for indice, linha in enumerate(list(obtido)):
        if linha is None:
            continue
        celas = list(linha)
        colunas = ["" if cela is None else str(cela).strip() for cela in celas]
        texto = " | ".join(colunas).strip(" |")
        if texto:
            rotulo = f"{nome_aba}#{indice + 1}"
            saida.append(
                TrechoExtraido(pagina_ou_aba=rotulo, texto=texto, origem=origem)
            )
    return saida


def extrair_docx_puro(caminho: Path) -> list[TrechoExtraido]:
    """Lê DOCX sem dependência, direto do XML interno."""
    try:
        with zipfile.ZipFile(str(caminho)) as pacote:
            bruto = pacote.read("word/document.xml")
    except (OSError, KeyError, zipfile.BadZipFile) as erro:
        raise ErroLeituraDocumento(f"docx ilegível: {caminho.name}") from erro
    etiqueta = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"
    textos = textos_de_xml(bruto, etiqueta)
    return dividir_texto("\n".join(textos), "documento", origem_base(caminho))


def extrair_xlsx_puro(caminho: Path) -> list[TrechoExtraido]:
    """Lê XLSX sem dependência, unindo textos compartilhados e abas."""
    try:
        with zipfile.ZipFile(str(caminho)) as pacote:
            nomes = pacote.namelist()
            nomes_abas = nomes_abas_puras(pacote, nomes)
            compartilhados = textos_compartilhados_puros(pacote, nomes)
            return abas_puras_para_trechos(
                pacote, nomes, nomes_abas, compartilhados, caminho
            )
    except (OSError, zipfile.BadZipFile) as erro:
        raise ErroLeituraDocumento(f"xlsx ilegível: {caminho.name}") from erro


def textos_de_xml(bruto: bytes, etiqueta: str) -> list[str]:
    """Extrai textos de etiqueta XML sem carregar esquema externo."""
    try:
        raiz = elemento_xml.fromstring(bruto)
    except elemento_xml.ParseError as erro:
        raise ErroLeituraDocumento("xml interno deformado") from erro
    return [no.text or "" for no in raiz.iter(etiqueta) if (no.text or "").strip()]


def nomes_abas_puras(pacote: zipfile.ZipFile, nomes: list[str]) -> list[str]:
    """Descobre abas pelo workbook ou cai para varredura de planilhas."""
    if "xl/workbook.xml" in nomes:
        achadas = _ler_nomes_workbook(pacote)
        if achadas:
            return achadas
    return [nome for nome in nomes if nome.startswith("xl/worksheets/")]


def textos_compartilhados_puros(pacote: zipfile.ZipFile, nomes: list[str]) -> list[str]:
    """Lê dicionário de textos compartilhados do XLSX."""
    if "xl/sharedStrings.xml" not in nomes:
        return []
    bruto = pacote.read("xl/sharedStrings.xml")
    etiqueta = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"
    return textos_de_xml(bruto, etiqueta)


def abas_puras_para_trechos(
    pacote: zipfile.ZipFile,
    nomes: list[str],
    nomes_abas: list[str],
    compartilhados: list[str],
    caminho: Path,
) -> list[TrechoExtraido]:
    """Monta um trecho por aba a partir de textos e valores puros."""
    arquivos = [nome for nome in nomes if nome.startswith("xl/worksheets/")]
    saida: list[TrechoExtraido] = []
    for posicao, arquivo_aba in enumerate(arquivos):
        rotulo = nomes_abas[posicao] if posicao < len(nomes_abas) else arquivo_aba
        bruto = pacote.read(arquivo_aba)
        texto = texto_aba_pura(bruto, compartilhados)
        if texto.strip():
            saida.append(
                TrechoExtraido(pagina_ou_aba=rotulo, texto=texto, origem=caminho.name)
            )
    return saida


def texto_aba_pura(bruto: bytes, compartilhados: list[str]) -> str:
    """Junta textos e índices numéricos da aba em linha legível."""
    try:
        raiz = elemento_xml.fromstring(bruto)
    except elemento_xml.ParseError as erro:
        raise ErroLeituraDocumento("aba xlsx deformada") from erro
    partes: list[str] = []
    for no in raiz.iter():
        conteudo = _conteudo_no(no)
        if not conteudo:
            continue
        etiqueta = _etiqueta_curta(no.tag)
        if etiqueta == "t":
            partes.append(conteudo)
        elif etiqueta == "v":
            partes.append(resolver_valor_puro(conteudo, compartilhados))
    return " | ".join(parte for parte in partes if parte)


def _conteudo_no(no: Any) -> str:
    """Devolve texto do nó sem espalhar `or` pelo leitor."""
    texto = no.text or ""
    return texto.strip()


def _etiqueta_curta(tag: str) -> str:
    """Tira namespace do XML para comparar só o nome local."""
    if "}" in tag:
        return tag.split("}")[-1]
    return tag


def resolver_valor_puro(valor: str, compartilhados: list[str]) -> str:
    """Troca índice de texto compartilhado pelo texto real."""
    if valor.isdigit():
        indice = int(valor)
        if indice < len(compartilhados):
            return compartilhados[indice]
    return valor


def _ler_nomes_workbook(pacote: zipfile.ZipFile) -> list[str]:
    """Lê nomes de abas do workbook, vazio se deformado."""
    try:
        bruto = pacote.read("xl/workbook.xml")
        raiz = elemento_xml.fromstring(bruto)
    except (KeyError, elemento_xml.ParseError):
        return []
    achadas = [no.attrib.get("name", "") for no in raiz.iter() if "sheet" in no.tag]
    return [nome for nome in achadas if nome]
